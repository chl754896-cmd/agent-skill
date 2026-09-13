"""Consistent, evidence-led V0.2 report rendering."""

import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import __version__

DIMENSIONS = {
    "goal_clarity": ("目标清晰度", "Goal clarity"), "capability_coverage": ("能力覆盖度", "Capability coverage"),
    "instruction_quality": ("指令质量", "Instruction quality"), "input_definition": ("输入定义", "Input definition"),
    "output_contract": ("输出契约", "Output contract"), "constraint_design": ("约束设计", "Constraint design"),
    "boundary_handling": ("边界处理", "Boundary handling"), "tool_workflow_design": ("工具与流程设计", "Tool and workflow design"),
    "robustness": ("鲁棒性", "Robustness"), "maintainability": ("可维护性", "Maintainability"),
    "missing_prerequisite": ("前置条件缺失", "Missing prerequisite"), "missing_validation": ("验证规则缺失", "Missing validation"),
    "unclear_priority": ("优先级不清晰", "Unclear priority"), "fallback": ("回退机制", "Fallback behavior"),
    "partial_input": ("部分输入", "Partial input"), "conflicting_instruction": ("冲突指令", "Conflicting instruction"),
}
RATINGS = {"excellent": ("优秀", "Excellent"), "strong": ("良好", "Strong"), "moderate": ("一般", "Moderate"), "weak": ("较弱", "Weak"), "critical": ("严重不足", "Critical")}
SEVERITIES = {"low": ("低", "Low"), "medium": ("中", "Medium"), "high": ("高", "High"), "critical": ("严重", "Critical")}

TEXT = {
    "zh-CN": {"title": "EvalForge 智能体技能评估与审计报告", "subtitle": "静态分析、动态评测与问题案例智能报告", "executive": "执行摘要", "profile": "技能档案", "capability": "能力地图", "capability_scores": "能力评分", "strengths": "优势分析", "weaknesses": "弱点与风险分析", "audit": "指令审计", "edge": "边界与鲁棒性分析", "dynamic": "动态评测结果", "badcase": "问题案例深度分析", "root": "根因分析", "recommendations": "P0 P1 P2 改进建议", "final": "最终评估", "metric": "指标", "value": "值", "skill": "技能名称", "score": "分数", "rating": "评级", "severity": "严重度", "evidence": "证据", "impact": "影响", "recommendation": "建议", "benefit": "预期收益", "effort": "预计工作量", "rationale": "优先级依据", "none": "未发现", "date": "评估日期", "version": "版本", "static": "静态得分", "dynamic_score": "动态得分", "composite": "综合得分", "maturity": "成熟度", "readiness": "生产就绪度", "critical_high": "严重或高风险发现数", "production": "生产建议", "why": "优势说明", "preserve": "应保留内容", "problem": "具体问题", "dimension": "所属维度", "case": "案例", "type": "类型", "reason": "原因", "source_note": "以下引用保留 SKILL.md 原文。", "ready": "具备条件上线", "conditional": "建议完成关键改进后上线", "not_ready": "暂不建议上线"},
    "en-US": {"title": "EvalForge Evaluation Report", "subtitle": "Agent Skill Analysis, Dynamic Evaluation, and Badcase Intelligence", "executive": "Executive Summary", "profile": "Skill Profile", "capability": "Capability Map", "capability_scores": "Capability Scores", "strengths": "Strength Analysis", "weaknesses": "Weaknesses and Risks", "audit": "Instruction Audit", "edge": "Edge Case and Robustness Analysis", "dynamic": "Dynamic Evaluation Results", "badcase": "Badcase Deep Analysis", "root": "Root Cause Analysis", "recommendations": "P0 P1 P2 Recommendations", "final": "Final Assessment", "metric": "Metric", "value": "Value", "skill": "Skill", "score": "Score", "rating": "Rating", "severity": "Severity", "evidence": "Evidence", "impact": "Impact", "recommendation": "Recommendation", "benefit": "Expected benefit", "effort": "Estimated effort", "rationale": "Priority rationale", "none": "None identified", "date": "Evaluation date", "version": "Version", "static": "Static score", "dynamic_score": "Dynamic score", "composite": "Composite score", "maturity": "Maturity level", "readiness": "Production readiness", "critical_high": "Critical or high findings", "production": "Production recommendation", "why": "Why this is strong", "preserve": "What to preserve", "problem": "Concrete problem", "dimension": "Affected dimension", "case": "Case", "type": "Type", "reason": "Reason", "source_note": "Source quotations below preserve the original SKILL.md wording.", "ready": "Ready for production", "conditional": "Ready after priority improvements", "not_ready": "Not ready for production"},
}


def _label(value, locale, mapping):
    return mapping.get(value, (value, value))[0 if locale == "zh-CN" else 1]


def _rating(score):
    return "excellent" if score >= 90 else "strong" if score >= 75 else "moderate" if score >= 60 else "weak" if score >= 40 else "critical"


def _bar(score):
    return "█" * round(score / 10) + "░" * (10 - round(score / 10))


def _rubric(item):
    return item.get("final_rubric") or item.get("rubric") or item.get("rule_rubric") or {"dimension_scores": {}, "overall_score": 0, "status": "fail"}


def _empty_static():
    return {"overall_score": 0, "overall_rating": "critical", "dimensions": {}, "strengths": [], "weaknesses": [], "instruction_audit": [], "edge_case_analysis": []}


def _priority(finding):
    """Reserve P0 for a critical correctness/safety blocker, not every high finding."""
    category, severity = finding.get("category", ""), finding.get("severity", "medium")
    if severity == "critical" and category in {"constraint_design", "robustness", "accuracy", "fallback", "missing_validation"}:
        return "P0"
    return "P1" if severity in {"high", "medium"} else "P2"


def _effort(finding):
    return "medium" if finding.get("severity") in {"high", "critical"} else "low"


def _recommendation(finding):
    priority = _priority(finding)
    rationale = {
        "P0": "Critical safety or correctness gap can block a reliable release.",
        "P1": "Important quality or robustness gap should be addressed before broader deployment.",
        "P2": "Maintainability or clarity optimization can follow the release-critical work.",
    }[priority]
    return {"priority": priority, "category": finding.get("category", ""), "severity": finding.get("severity", "medium"), "title": finding.get("title", "Quality improvement"), "problem": finding.get("conclusion", finding.get("reason", "A quality gap was observed.")), "evidence": "; ".join(item.get("quote_or_summary", "") for item in finding.get("evidence", []) if isinstance(item, dict)) or finding.get("reason", "No direct evidence recorded."), "impact": finding.get("impact", "The Agent may behave inconsistently."), "suggested_change": finding.get("recommendation", "Add a concrete validation rule and regression case."), "expected_benefit": finding.get("expected_benefit", "The Skill becomes easier to execute and verify consistently."), "priority_rationale": rationale, "estimated_effort": _effort(finding)}


def _maintenance_recommendation(profile):
    evidence = "; ".join(profile.get("constraints", []) + profile.get("prohibitions", [])) or "The Skill already has an explicit structure."
    return {"priority": "P2", "category": "constraint_design", "severity": "low", "title": "Maintain the explicit constraint contract", "problem": "The current safeguards are valuable but can drift when the Skill is edited.", "evidence": evidence, "impact": "A future edit could weaken the no-invention or output-format expectations.", "suggested_change": "Keep the existing constraints in regression review and update their examples with each material Skill change.", "expected_benefit": "Preserves predictable behavior while keeping the document maintainable.", "priority_rationale": "This is a maintainability optimization, not a release blocker.", "estimated_effort": "low"}


def _scorecard(static, dynamic_score, findings):
    has_static = bool(static.get("dimensions"))
    static_score = static.get("overall_score", 0) if has_static else dynamic_score
    composite = round(static_score * 0.55 + dynamic_score * 0.45) if has_static else dynamic_score
    rating = _rating(composite)
    p0 = any(_priority(item) == "P0" for item in findings)
    if composite >= 85 and not p0:
        maturity, readiness = "managed-ready", "ready"
    elif composite >= 70 and not p0:
        maturity, readiness = "conditionally-ready", "conditional"
    else:
        maturity, readiness = "early-stage", "not_ready"
    return {"static_score": static_score, "dynamic_score": dynamic_score, "composite_score": composite, "rating": rating, "maturity_level": maturity, "production_readiness": readiness, "static_weight": 0.55 if has_static else 0.0, "dynamic_weight": 0.45 if has_static else 1.0, "critical_high_findings": sum(item.get("severity") in {"critical", "high"} for item in findings)}


def build_report(evaluation):
    """Build one consistent scorecard used by JSON, Markdown, and DOCX."""
    results = evaluation.get("results", [])
    rubrics = [_rubric(item) for item in results]
    statuses, values = Counter(item["status"] for item in rubrics), {}
    for rubric in rubrics:
        for name, score in rubric.get("dimension_scores", {}).items(): values.setdefault(name, []).append(score)
    dynamic_score = round(sum(item["overall_score"] for item in rubrics) / len(rubrics)) if rubrics else 0
    dynamic_summary = {"total_cases": len(results), "pass": statuses["pass"], "warning": statuses["warning"], "fail": statuses["fail"], "pass_rate": round(statuses["pass"] / len(results) * 100, 1) if results else 0, "overall_score": dynamic_score, "status": "pass" if not statuses["fail"] and not statuses["warning"] else "warning" if not statuses["fail"] else "fail"}
    static = evaluation.get("static_analysis") or _empty_static()
    profile = evaluation.get("skill_profile") or {"name": evaluation.get("skill", "Unknown Skill"), "purpose": "", "target_scenarios": [], "inputs": [], "outputs": [], "tools": [], "constraints": [], "prohibitions": [], "workflow_steps": [], "failure_handling": [], "dependencies": []}
    badcases = [item["badcase"] for item in results if item.get("badcase")]
    all_findings = static.get("weaknesses", []) + static.get("instruction_audit", []) + static.get("edge_case_analysis", []) + badcases
    scorecard = _scorecard(static, dynamic_score, all_findings)
    recommendations = {"P0": [], "P1": [], "P2": []}
    for finding in all_findings:
        item = _recommendation(finding); recommendations[item["priority"]].append(item)
    recommendations["P2"].append(_maintenance_recommendation(profile))
    rationale = {
        "P0": "No critical correctness or safety release blocker was found." if not recommendations["P0"] else "Items below are critical correctness or safety blockers.",
        "P1": "Items below improve robustness, input handling, or execution quality before wider rollout." if recommendations["P1"] else "No important quality gap was found in the available evidence.",
        "P2": "Items below preserve clarity and maintainability after higher-priority work.",
    }
    root_causes = [{"case_id": item["case_id"], "badcase_type": item["type"], "reason": item["reason"], "root_cause": item["root_cause"], "recommendation": item["recommendation"]} for item in badcases]
    narrative = "Dynamic responses passed the supplied cases, but the static review found unresolved design gaps. Complete the P1 items before broad production deployment." if scorecard["production_readiness"] == "conditional" else "The combined static and dynamic evidence supports production use with the documented controls." if scorecard["production_readiness"] == "ready" else "The combined evidence is insufficient for reliable production use; resolve the P0 and P1 items first."
    final = {**scorecard, "narrative": narrative, "key_strengths": [item.get("title", "") for item in static.get("strengths", [])[:3]], "key_risks": [item.get("title", "") for item in all_findings if item.get("severity") in {"critical", "high", "medium"}][:3], "next_steps": [item["suggested_change"] for priority in ("P0", "P1", "P2") for item in recommendations[priority][:1]]}
    dynamic = {"summary": dynamic_summary, "capability_scores": {name: round(sum(scores) / len(scores)) for name, scores in values.items()}, "case_results": results, "badcases": badcases, "badcase_distribution": dict(Counter(item["type"] for item in badcases))}
    # summary is retained for V0.1 consumers, but now deliberately exposes the composite score.
    summary = {**dynamic_summary, **scorecard, "overall_score": scorecard["composite_score"]}
    return {"meta": {"evalforge_version": __version__, "skill_path": evaluation.get("skill_path", ""), "generated_at": datetime.now(timezone.utc).isoformat()}, "scorecard": scorecard, "skill_profile": profile, "static_analysis": static, "dynamic_evaluation": dynamic, "root_cause_analysis": root_causes, "recommendations": recommendations, "recommendation_rationale": rationale, "final_assessment": final, "summary": summary, "capability_scores": dynamic["capability_scores"], "failed_cases": badcases, "badcase_distribution": dynamic["badcase_distribution"]}


def _localized_finding(item, locale, strength=False):
    category = item.get("category", "")
    dimension = _label(category, locale, DIMENSIONS)
    severity = _label(item.get("severity", "medium"), locale, SEVERITIES)
    profile = {
        "goal_clarity": (("目标表述直接说明了简历审核的范围和预期价值。", "明确的目标帮助 Agent 优先处理表达与结构问题。", "继续将目标保持在文档开头，并在扩展功能时同步更新。"), ("The purpose states the scope and intended value of resume review directly.", "A clear goal helps an Agent prioritize expression and structure issues.", "Keep the goal near the top and update it whenever the Skill scope changes.")),
        "capability_coverage": (("输出列出了总体评价、优势、问题和修改建议，覆盖了审核任务的核心交付物。", "明确的交付物使结果更便于用户核对和消费。", "保留这四类输出，并在新增能力时补充对应验收项。"), ("The listed outputs cover the core deliverables: assessment, strengths, issues, and actionable improvements.", "A visible deliverable set makes results easier for users to verify and use.", "Preserve these output groups and add acceptance criteria when capabilities expand.")),
        "instruction_quality": (("目标、约束和示例共同提供了可执行的审核方向。", "这些指令减少了 Agent 对任务目的的主观猜测。", "保留示例与约束之间的一致性，并为新增规则补充示例。"), ("The purpose, constraints, and example provide usable review direction together.", "These instructions reduce subjective guessing about the task intent.", "Keep examples aligned with constraints and add examples for material new rules.")),
        "input_definition": (("输入说明明确了简历可能包含的主要信息范围。", "Agent 可以据此识别可分析的信息与缺失信息。", "保留输入范围，并在后续补充必填与可选字段。"), ("The input section identifies the main information a resume may contain.", "An Agent can use it to distinguish analysable information from missing information.", "Preserve the scope and later distinguish required from optional fields.")),
        "output_contract": (("输出以清晰列表规定了审核结果的组成。", "稳定的输出契约有助于下游使用者比较不同审核结果。", "保留列表结构，并为每项增加最小内容要求。"), ("The output list defines the components of a review result clearly.", "A stable output contract helps downstream users compare review results.", "Preserve the list and add minimum content requirements for each item.")),
        "constraint_design": (("禁止虚构经历和要求分点输出等规则为审核设置了明确边界。", "这些边界降低了无依据补充信息和格式漂移的风险。", "保持禁止项可见，并将其纳入后续回归用例。"), ("The no-invention rule and list-format requirements set clear boundaries for review behavior.", "These boundaries reduce unsupported additions and format drift.", "Keep prohibitions visible and retain them in future regression cases.")),
        "maintainability": (("文档使用了目标、输入、输出、约束和示例等稳定章节。", "结构化章节使规则更新时更容易定位影响范围。", "保留章节结构，并在新增流程或回退规则时使用同一层级。"), ("The document uses stable sections for purpose, inputs, outputs, constraints, and examples.", "Structured sections make the impact of rule changes easier to locate.", "Preserve the hierarchy and use it for future workflow or fallback rules.")),
        "boundary_handling": (("Skill 未说明面对缺失、格式错误、冲突或不支持输入时应如何处理。", "不同 Agent 可能自行选择澄清、拒绝或继续推断，导致行为不一致。", "新增失败处理章节，分别规定澄清、拒绝、降级输出和停止条件。"), ("The Skill does not state how to handle missing, malformed, conflicting, or unsupported input.", "Different Agents may choose to clarify, refuse, or infer, producing inconsistent behavior.", "Add a failure-handling section defining clarification, refusal, degraded output, and stop conditions.")),
        "tool_workflow_design": (("Skill 没有给出审核步骤、规则优先级或工具使用顺序。", "相同输入可能因执行顺序不同而得到不一致的审核深度。", "增加有序流程：校验输入、识别缺失信息、执行审核、核对禁止项、生成分点建议。"), ("The Skill does not define review steps, rule priority, or tool ordering.", "The same input can receive different review depth when execution order changes.", "Add an ordered flow: validate input, identify gaps, review content, check prohibitions, then generate bullet recommendations.")),
        "robustness": (("现有约束未配套异常输入的验证与回退规则。", "当信息不完整或指令冲突时，Agent 可能遗漏澄清或给出越界建议。", "为缺失信息、冲突指令和不支持格式分别定义验证规则与安全回退。"), ("Existing constraints are not paired with validation and fallback rules for exceptional input.", "When information is incomplete or instructions conflict, an Agent may omit clarification or make out-of-bound recommendations.", "Define validation and safe fallback rules for missing information, conflicts, and unsupported formats.")),
        "unclear_priority": (("未定义流程顺序或规则优先级。", "规则之间发生冲突时，Agent 无法稳定判断应先遵守哪一条。", "声明约束和禁止项高于用户附加指令，并给出执行顺序。"), ("No workflow order or rule priority is defined.", "When rules conflict, an Agent cannot consistently decide which one governs.", "State that constraints and prohibitions outrank user additions and provide an execution order.")),
        "fallback": (("没有针对异常输入的回退机制。", "缺失或不支持的信息可能导致不一致输出或无依据推断。", "定义澄清问题、拒绝条件和最小可用输出。"), ("No fallback is defined for exceptional input.", "Missing or unsupported information can lead to inconsistent output or unsupported inference.", "Define clarification prompts, refusal conditions, and a minimum viable output.")),
        "partial_input": (("输入范围说明了多类简历信息，但没有规定缺少关键字段后的处理方式。", "部分输入可能被误当成完整信息，降低审核建议的可靠性。", "增加缺少项目规模、成果指标或职责时的澄清用例和处理规则。"), ("The input scope lists several resume elements but does not define handling when key fields are absent.", "Partial input may be treated as complete, reducing the reliability of recommendations.", "Add clarification rules and cases for missing project scale, outcomes, or responsibilities.")),
        "conflicting_instruction": (("约束存在，但未明确说明它们如何优先于冲突的用户指令。", "冲突请求可能绕过禁止虚构等重要边界。", "在流程中写明禁止项和约束的优先级，并加入冲突指令回归测试。"), ("Constraints exist, but their priority over conflicting user instructions is not stated.", "A conflicting request could bypass important boundaries such as the no-invention rule.", "State the priority of prohibitions and constraints in the flow and add conflict-regression tests.")),
    }
    content = profile.get(category)
    if content:
        conclusion, impact, recommendation = content[0 if locale == "zh-CN" else 1]
    else:
        conclusion, impact, recommendation = item.get("conclusion", "Missing documented behavior."), item.get("impact", "Behavior may be inconsistent."), item.get("recommendation", "Add explicit guidance.")
    title = _label(category, locale, DIMENSIONS) if category in DIMENSIONS else item.get("title", category)
    evidence = item.get("evidence", [])
    return {"title": title, "dimension": dimension, "severity": severity, "conclusion": conclusion, "evidence": evidence, "impact": impact, "recommendation": recommendation, "why": conclusion if strength else "", "preserve": recommendation if strength else ""}


def _profile_rows(profile, locale):
    labels = {"name": ("名称", "Name"), "purpose": ("目的", "Purpose"), "target_scenarios": ("适用场景", "Target scenarios"), "inputs": ("输入", "Inputs"), "outputs": ("输出", "Outputs"), "tools": ("工具", "Tools"), "constraints": ("约束", "Constraints"), "prohibitions": ("禁止项", "Prohibitions"), "workflow_steps": ("流程步骤", "Workflow steps"), "failure_handling": ("失败处理", "Failure handling"), "dependencies": ("依赖", "Dependencies")}
    result = []
    for key, label in labels.items():
        value = profile.get(key, [])
        value = "; ".join(value) if isinstance(value, list) else value
        if not value:
            value = "未提供" if locale == "zh-CN" else "Not provided"
        elif locale == "en-US" and any("\u4e00" <= character <= "\u9fff" for character in value):
            value = f"Declared in SKILL.md: “{value}”"
        result.append((_label(key, locale, labels), value))
    return result


def _dynamic_label(key, locale):
    labels = {"total_cases": ("Case 总数", "Total cases"), "pass": ("通过", "Pass"), "warning": ("警告", "Warning"), "fail": ("失败", "Fail"), "pass_rate": ("通过率", "Pass rate"), "overall_score": ("动态总分", "Dynamic overall score"), "status": ("动态状态", "Dynamic status")}
    return _label(key, locale, labels)


def _localized_recommendation(item, locale):
    if locale == "en-US":
        return item
    finding = _localized_finding({"category": item.get("category", ""), "severity": item.get("severity", "medium"), "evidence": []}, locale)
    evidence = item.get("evidence", "")
    missing_evidence = {"boundary_handling": "SKILL.md 中未提供失败处理章节。", "tool_workflow_design": "SKILL.md 中未提供流程或工具使用顺序。", "robustness": "SKILL.md 中未提供异常输入的验证与回退规则。", "unclear_priority": "SKILL.md 中未提供规则优先级或执行顺序。", "fallback": "SKILL.md 中未提供失败处理或回退机制。"}
    if evidence.startswith("The ") or evidence.startswith("No "):
        evidence = missing_evidence.get(item.get("category"), evidence)
    benefits = {"boundary_handling": "异常输入将得到一致、可预测的处理。", "tool_workflow_design": "相同输入可获得更稳定的审核过程与输出。", "robustness": "减少缺失信息和冲突指令导致的越界建议。", "unclear_priority": "规则冲突时可稳定执行高优先级约束。", "fallback": "面对异常输入时可避免无依据推断。", "partial_input": "部分简历信息将触发澄清，而非被当成完整资料。", "conflicting_instruction": "冲突请求不会绕过禁止虚构等核心边界。", "constraint_design": "后续迭代可持续保留关键约束。"}
    return {**item, "title": finding["title"], "problem": finding["conclusion"], "evidence": evidence, "impact": finding["impact"], "suggested_change": finding["recommendation"], "expected_benefit": benefits.get(item.get("category"), "提高 Skill 的可预测性与可验证性。")}


def markdown_report(report, locale="en-US"):
    t, scorecard, static, dynamic = TEXT[locale], report["scorecard"], report["static_analysis"], report["dynamic_evaluation"]
    lines = [f"# {t['title']}", "", t["subtitle"], "", f"- **{t['skill']}**: {report['skill_profile'].get('name', '')}", f"- **{t['version']}**: {report['meta']['evalforge_version']}", f"- **{t['date']}**: {report['meta']['generated_at']}"]
    lines += ["", f"## {t['executive']}", "", f"| {t['metric']} | {t['value']} |", "| --- | --- |", f"| {t['static']} | {scorecard['static_score']} |", f"| {t['dynamic_score']} | {scorecard['dynamic_score']} |", f"| {t['composite']} | **{scorecard['composite_score']}** |", f"| {t['rating']} | {_label(scorecard['rating'], locale, RATINGS)} |", f"| {t['maturity']} | {_maturity(scorecard['maturity_level'], locale)} |", f"| {t['readiness']} | {t[scorecard['production_readiness']]} |", f"| {t['critical_high']} | {scorecard['critical_high_findings']} |", "", f"**{t['production']}**: {_narrative(report['final_assessment']['narrative'], locale)}", "", f"**{t['strengths']}**: " + "; ".join(_label(item.get('category', ''), locale, DIMENSIONS) for item in static.get('strengths', [])[:3]), f"**{t['weaknesses']}**: " + "; ".join(_label(item.get('category', ''), locale, DIMENSIONS) for item in (static.get('weaknesses', []) + static.get('instruction_audit', []) + static.get('edge_case_analysis', []))[:3])]
    lines += ["", f"## {t['profile']}", "", t["source_note"], "", f"| {t['metric']} | {t['value']} |", "| --- | --- |"] + [f"| {key} | {value} |" for key, value in _profile_rows(report["skill_profile"], locale)]
    lines += ["", f"## {t['capability']}", "", f"| {t['dimension']} | {t['score']} | {t['rating']} |", "| --- | ---: | --- |"] + [f"| {_label(key, locale, DIMENSIONS)} | {_bar(value['score'])} {value['score']} | {_label(value['rating'], locale, RATINGS)} |" for key, value in static.get("dimensions", {}).items()]
    for heading, key, strength in ((t["strengths"], "strengths", True), (t["weaknesses"], "weaknesses", False), (t["audit"], "instruction_audit", False), (t["edge"], "edge_case_analysis", False)):
        lines += ["", f"## {heading}", ""]
        items = static.get(key, [])
        if not items: lines.append(t["none"])
        for item in items:
            found = _localized_finding(item, locale, strength)
            lines += [f"### {found['title']}", f"- **{t['dimension']}**: {found['dimension']}", f"- **{t['severity']}**: {found['severity']}"]
            if strength: lines += [f"- **{t['why']}**: {found['why']}"]
            else: lines += [f"- **{t['problem']}**: {found['conclusion']}"]
            for evidence in found["evidence"]:
                if isinstance(evidence, dict): lines.append(f"- **{t['evidence']}**: {evidence.get('quote_or_summary', '')} ({evidence.get('section', 'SKILL.md')})")
            lines += [f"- **{t['impact']}**: {found['impact']}", f"- **{t['preserve'] if strength else t['recommendation']}**: {found['preserve'] if strength else found['recommendation']}", ""]
    lines += [f"## {t['dynamic']}", "", f"| {t['metric']} | {t['value']} |", "| --- | --- |"] + [f"| {_dynamic_label(key, locale)} | {value} |" for key, value in dynamic['summary'].items()]
    lines += ["", f"## {t['capability_scores']}", "", f"| {t['dimension']} | {t['score']} |", "| --- | ---: |"] + [f"| {_label(key, locale, DIMENSIONS)} | {value} |" for key, value in dynamic['capability_scores'].items()]
    lines += ["", f"## {t['badcase']}", "", f"| {t['case']} | {t['type']} | {t['severity']} | {t['reason']} |", "| --- | --- | --- | --- |"]
    lines += [f"| {item['case_id']} | {item['type']} | {_label(item.get('severity', 'medium'), locale, SEVERITIES)} | {item.get('reason', '')} |" for item in dynamic['badcases']] or [f"| - | {t['none']} | - | - |"]
    lines += ["", f"## {t['root']}", ""]
    if report['root_cause_analysis']:
        for item in report['root_cause_analysis']: lines += [f"- **{item['case_id']}**: {item['root_cause']} — {item['recommendation']}"]
    else: lines.append(t['none'])
    lines += ["", f"## {t['recommendations']}", ""]
    for priority in ("P0", "P1", "P2"):
        rationale = _recommendation_rationale(priority, locale) if report['recommendations'][priority] else ("当前未发现 P0 发布阻塞项。" if locale == 'zh-CN' and priority == 'P0' else report['recommendation_rationale'][priority])
        lines += [f"### {priority}", rationale]
        for raw_item in report['recommendations'][priority]:
            item = _localized_recommendation(raw_item, locale)
            lines += [f"- **{item['title']}**: {item['problem']}", f"  - {t['evidence']}: {item['evidence']}", f"  - {t['impact']}: {item['impact']}", f"  - {t['recommendation']}: {item['suggested_change']}", f"  - {t['benefit']}: {item['expected_benefit']}", f"  - {t['rationale']}: {_recommendation_rationale(item['priority'], locale)}", f"  - {t['effort']}: {_effort_label(item['estimated_effort'], locale)}"]
    final = report['final_assessment']
    lines += ["", f"## {t['final']}", "", f"- **{t['composite']}**: {final['composite_score']} ({_label(final['rating'], locale, RATINGS)})", f"- **{t['maturity']}**: {_maturity(final['maturity_level'], locale)}", f"- **{t['readiness']}**: {t[final['production_readiness']]}", f"- **{t['production']}**: {_narrative(final['narrative'], locale)}"]
    return "\n".join(lines) + "\n"


def _maturity(value, locale):
    labels = {"managed-ready": ("可管理上线", "Managed-ready"), "conditionally-ready": ("有条件上线", "Conditionally ready"), "early-stage": ("早期阶段", "Early stage")}
    return _label(value, locale, labels)


def _narrative(value, locale):
    if locale == "en-US": return value
    return {"Dynamic responses passed the supplied cases, but the static review found unresolved design gaps. Complete the P1 items before broad production deployment.": "提供的模型响应已通过动态 Case，但静态审计仍发现未解决的设计缺口；建议完成 P1 项后再扩大生产使用范围。", "The combined static and dynamic evidence supports production use with the documented controls.": "静态与动态证据共同支持在已记录控制措施下投入生产使用。", "The combined evidence is insufficient for reliable production use; resolve the P0 and P1 items first.": "当前静态与动态证据不足以支持可靠生产使用；应先完成 P0 和 P1 项。"}.get(value, value)


def _recommendation_rationale(priority, locale):
    values = {"P0": ("属于关键正确性或安全发布阻塞项。", "This is a critical correctness or safety release blocker."), "P1": ("该项会显著影响鲁棒性或交付质量，应在扩大上线前完成。", "This materially affects robustness or delivery quality and should be completed before broader rollout."), "P2": ("该项用于提升可维护性或长期清晰度，可在高优先级工作后安排。", "This improves maintainability or long-term clarity and can follow higher-priority work.")}
    return _label(priority, locale, values)


def _effort_label(value, locale):
    return _label(value, locale, {"low": ("低", "Low"), "medium": ("中", "Medium"), "high": ("高", "High")})


def _shade(cell, color):
    node = OxmlElement("w:shd"); node.set(qn("w:fill"), color); cell._tc.get_or_add_tcPr().append(node)


def _repeat(row):
    node = OxmlElement("w:tblHeader"); node.set(qn("w:val"), "true"); row._tr.get_or_add_trPr().append(node)


def _table(document, headers, rows, badge_column=None):
    table = document.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"; _repeat(table.rows[0])
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value); _shade(cell, "1F4E78"); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs: run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.text = str(value); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2: _shade(cell, "EAF3F8")
            if badge_column == index:
                color = "C6E0B4" if str(value) in {"P2", "优秀", "Excellent", "良好", "Strong"} else "FFE699" if str(value) in {"P1", "一般", "Moderate", "中", "Medium"} else "F4CCCC"
                _shade(cell, color)
    document.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def _page_number(paragraph):
    paragraph.add_run("Page "); field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); paragraph._p.append(field)


def _heading(document, title, major=False):
    if major: document.add_page_break()
    document.add_heading(title, level=1)


def _docx_finding(document, item, locale, strength=False):
    t, found = TEXT[locale], _localized_finding(item, locale, strength)
    document.add_heading(found['title'], level=2)
    rows = [(t['dimension'], found['dimension']), (t['severity'], found['severity']), (t['why'] if strength else t['problem'], found['why'] if strength else found['conclusion'])]
    for evidence in found['evidence']:
        if isinstance(evidence, dict): rows.append((t['evidence'], f"{evidence.get('quote_or_summary', '')} ({evidence.get('section', 'SKILL.md')})"))
    rows += [(t['impact'], found['impact']), (t['preserve'] if strength else t['recommendation'], found['preserve'] if strength else found['recommendation'])]
    _table(document, [t['metric'], t['value']], rows)


def docx_report(report, path, locale="en-US"):
    """Render a compact, locale-specific professional Word report."""
    t, scorecard, static, dynamic = TEXT[locale], report['scorecard'], report['static_analysis'], report['dynamic_evaluation']
    document = Document(); section = document.sections[0]
    for side in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'): setattr(section, side, Inches(.72))
    normal = document.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(10.5); normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for name, size in (('Title', 24), ('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 11)):
        style = document.styles[name]; style.font.name = 'Arial'; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0); style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    section.header.paragraphs[0].text = 'EvalForge'; section.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; _page_number(footer)
    title = document.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.add_run(t['title'])
    subtitle = document.add_paragraph(t['subtitle']); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER; subtitle.paragraph_format.space_after = Pt(18)
    _table(document, [t['metric'], t['value']], [(t['skill'], report['skill_profile'].get('name', '')), (t['date'], report['meta']['generated_at']), (t['version'], report['meta']['evalforge_version']), (t['composite'], scorecard['composite_score']), (t['rating'], _label(scorecard['rating'], locale, RATINGS)), (t['readiness'], t[scorecard['production_readiness']])], badge_column=1)
    _heading(document, t['executive'])
    _table(document, [t['metric'], t['value']], [(t['static'], scorecard['static_score']), (t['dynamic_score'], scorecard['dynamic_score']), (t['composite'], scorecard['composite_score']), (t['maturity'], _maturity(scorecard['maturity_level'], locale)), (t['readiness'], t[scorecard['production_readiness']]), (t['critical_high'], scorecard['critical_high_findings'])], badge_column=1)
    document.add_paragraph(f"{t['production']}: {_narrative(report['final_assessment']['narrative'], locale)}")
    _heading(document, t['profile'], major=True); document.add_paragraph(t['source_note']); _table(document, [t['metric'], t['value']], _profile_rows(report['skill_profile'], locale))
    _heading(document, t['capability']); _table(document, [t['dimension'], t['score'], t['rating']], [(_label(key, locale, DIMENSIONS), f"{_bar(value['score'])} {value['score']}", _label(value['rating'], locale, RATINGS)) for key, value in static.get('dimensions', {}).items()], badge_column=2)
    for heading, key, strength in ((t['strengths'], 'strengths', True), (t['weaknesses'], 'weaknesses', False), (t['audit'], 'instruction_audit', False), (t['edge'], 'edge_case_analysis', False)):
        _heading(document, heading, major=heading in {t['weaknesses'], t['dynamic']})
        if not static.get(key): document.add_paragraph(t['none'])
        for item in static.get(key, []): _docx_finding(document, item, locale, strength)
    _heading(document, t['dynamic'], major=True); _table(document, [t['metric'], t['value']], [(_dynamic_label(key, locale), value) for key, value in dynamic['summary'].items()]); _table(document, [t['dimension'], t['score']], [(_label(key, locale, DIMENSIONS), value) for key, value in dynamic['capability_scores'].items()])
    _heading(document, t['badcase']); _table(document, [t['case'], t['type'], t['severity'], t['reason']], [(item['case_id'], item['type'], _label(item.get('severity', 'medium'), locale, SEVERITIES), item.get('reason', '')) for item in dynamic['badcases']] or [('-', t['none'], '-', '-')])
    _heading(document, t['root']);
    if report['root_cause_analysis']: _table(document, [t['case'], t['reason'], t['recommendation']], [(item['case_id'], item['root_cause'], item['recommendation']) for item in report['root_cause_analysis']])
    else: document.add_paragraph(t['none'])
    _heading(document, t['recommendations'], major=True)
    for priority in ('P0', 'P1', 'P2'):
        document.add_heading(priority, level=2); document.add_paragraph(_recommendation_rationale(priority, locale) if report['recommendations'][priority] else ("当前未发现 P0 发布阻塞项。" if locale == 'zh-CN' and priority == 'P0' else report['recommendation_rationale'][priority]))
        if not report['recommendations'][priority]: document.add_paragraph(t['none'])
        for raw_item in report['recommendations'][priority]:
            item = _localized_recommendation(raw_item, locale)
            _table(document, [t['metric'], t['value']], [(t['rationale'], _recommendation_rationale(priority, locale)), (t['problem'], item['problem']), (t['evidence'], item['evidence']), (t['impact'], item['impact']), (t['recommendation'], item['suggested_change']), (t['benefit'], item['expected_benefit']), (t['effort'], _effort_label(item['estimated_effort'], locale))], badge_column=1)
    _heading(document, t['final'], major=True); final = report['final_assessment']; _table(document, [t['metric'], t['value']], [(t['static'], final['static_score']), (t['dynamic_score'], final['dynamic_score']), (t['composite'], final['composite_score']), (t['rating'], _label(final['rating'], locale, RATINGS)), (t['maturity'], _maturity(final['maturity_level'], locale)), (t['readiness'], t[final['production_readiness']])], badge_column=1); document.add_heading('Evaluation Conclusion' if locale == 'en-US' else '评估结论', level=2); document.add_paragraph(_narrative(final['narrative'], locale)); document.save(path)


def write_reports(evaluation, markdown_path='eval_report.md', json_path='eval_report.json', docx_path='eval_report.docx'):
    report = build_report(evaluation); base = Path(json_path).parent
    zh_md, en_md, zh_docx, en_docx = base / 'eval_report_zh-CN.md', base / 'eval_report_en-US.md', base / 'eval_report_zh-CN.docx', base / 'eval_report_en-US.docx'
    zh_md.write_text(markdown_report(report, 'zh-CN'), encoding='utf-8'); en_md.write_text(markdown_report(report, 'en-US'), encoding='utf-8'); Path(json_path).write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    docx_report(report, zh_docx, 'zh-CN'); docx_report(report, en_docx, 'en-US')
    Path(markdown_path).write_text(markdown_report(report, 'en-US'), encoding='utf-8')
    if Path(docx_path) != en_docx: docx_report(report, docx_path, 'en-US')
    return report
